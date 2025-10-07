# This is the main python file that will be used to pull in the information/code from all other files in my Stuent Feedback Project


from dictionaries.student_names import student
#print(student_names) #added the print to seach section so I could test as I was going along.

from dictionaries.performance_categories import performance_categories
#print(performance_categories)

from dictionaries.score_descriptions import score_descriptions
#print(score_desciptions)

import os

from docx import Document



print("Welcome to the Student Feedback Report\n")
input("Press Enter...\n")

# This part will show the performance categoriess
print("As the Tutor you will be marking your students on the following\n")
for category in performance_categories:
    print(f"- {category}")
print() # This is a blank line for when code is run
input("Press Enter...\n")

# This is my scoring key    
print("Using the following scoring key, you need to mark each category 1 - 4 for each student\n")
for key, value in score_descriptions.items():
    print(f"{key} = {value}")
    
print() # This is a blank line for when code is run    

for student_name in student:
    print(f"\nEnter scores for {student_name}:")
    for category in performance_categories:
        while True:
            try:
                score = int(input(f"- {category} (1-4): "))
                if score in [1, 2, 3, 4]:
                    student[student_name][category] = score  # store the score
                    break  # exit the while loop
                else:
                    print("Please enter a number between 1 and 4.")
            except ValueError:
                print("Invalid input. Please enter a number between 1 and 4.")

print(f"\nScores for {student} recorded successfully!")

# You can change this path for where the _Perfromance_Score.txt file will be saved
output_folder = r"C:\Artemis\Python\Project\student_feedback_plm\student_reports"
os.makedirs(output_folder, exist_ok=True)

for student_name, scores in student.items():
     # You can change the naming convention here: e.g from _PerformanceScore.txt to _Performance_Score.txt
    filename = os.path.join(output_folder, f"{student_name}_Performance_Score.txt")
    
    with open(filename, "w") as file:
        file.write(f"Student Performance Score for {student_name}\n")
        file.write("="*40 + "\n")
        for category, score in scores.items():
            file.write(f"{category}: {score}\n")
    
    print(f"Report created for {student_name}: {filename}")



# Feedback word doc using my template
# You can change this path for where the _feedback.docx file will be saved
output_folder = r"C:\Artemis\Python\Project\student_feedback_plm\feedback"
os.makedirs(output_folder, exist_ok=True)

# This is the path to the template that is used for the final _feedback.docx file
template_path = r"C:\Artemis\Python\Project\student_feedback_plm\templates\feedback\template.docx"

# This generate the feedback doc for each
for student_name, scores in student.items():
    # This loads the template
    doc = Document(template_path)

    # This replaces the  placeholders in the template
    for paragraph in doc.paragraphs:
        if "{{STUDENT_NAME}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{STUDENT_NAME}}", student_name)
        if "{{SCORES}}" in paragraph.text:
            
            scores_text = "\n".join(f"{category}: {score}" for category, score in scores.items())
            paragraph.text = paragraph.text.replace("{{SCORES}}", scores_text)

    # You can change the naming convention here: e.g from _feed_back.docx to _feedback.docx
    docx_filename = os.path.join(output_folder, f"{student_name}_feedback.docx")
    doc.save(docx_filename)

    # This automatically opens the word doc
    os.startfile(docx_filename)

    print(f"Word report created and opened for {student_name}: {docx_filename}")